#!/usr/bin/env python3
"""
Excel с карточками «Готовых решений» -> public/data/apps/<id>.json

Формат листа (как в «ЛМ_основа_работа.xlsx»): категории идут блоками по 3 столбца.
  строка 1      — «эмодзи Название»
  строки 2-3    — подпись («→ тема, тема, тема»)
  строки 4...   — посты: текст ячейки = заголовок, ГИПЕРССЫЛКА ячейки = ссылка на пост
  последняя     — «Боль: «цитата»»

Использование:
  python3 scripts/xlsx-to-app-content.py файл.xlsx public/data/apps/jobs.json --app jobs --channel "Работа не рабство"
"""
import argparse, json, re, sys
import openpyxl

TRANSLIT = dict(zip("абвгдеёжзийклмнопрстуфхцчшщъыьэюя",
                    ["a","b","v","g","d","e","e","zh","z","i","y","k","l","m","n","o","p","r","s","t","u","f","h","c","ch","sh","sch","","y","","e","yu","ya"]))

def slugify(title):
    s = "".join(TRANSLIT.get(ch, ch) for ch in title.lower())
    return re.sub(r"[^a-z0-9]+", "_", s).strip("_") or "category"

def clean_url(u):
    # убираем метки вроде ?utm_source=chatgpt.com
    return re.sub(r"[?&]utm_[^&]*", "", u).rstrip("?&") if u else u

def convert(path, app_id, channel):
    ws = openpyxl.load_workbook(path).active
    anchor = {}
    for r in ws.merged_cells.ranges:
        for row in range(r.min_row, r.max_row + 1):
            for col in range(r.min_col, r.max_col + 1):
                anchor[(row, col)] = (r.min_row, r.min_col)
    cell = lambda row, col: ws.cell(*anchor.get((row, col), (row, col)))

    cats, problems, used = [], [], set()
    for c0 in range(1, ws.max_column + 1, 3):
        head = (ws.cell(1, c0).value or "").strip()
        if not head:
            continue
        m = re.match(r"^(\S+)\s+(.*)$", head)
        emoji, title = (m.group(1), m.group(2)) if m else ("", head)
        slug = slugify(title)
        while slug in used:
            slug += "_2"
        used.add(slug)
        sub = (cell(2, c0).value or "").lstrip("→ ").strip()
        posts, pain, seen = [], None, set()
        for row in range(4, ws.max_row + 1):
            cc = cell(row, c0)
            if (cc.row, cc.column) in seen or not cc.value:
                continue
            seen.add((cc.row, cc.column))
            text = str(cc.value).strip()
            if text.startswith("Боль:"):
                pain = text[len("Боль:"):].strip()
                continue
            url = clean_url(cc.hyperlink.target) if cc.hyperlink else None
            if not url:
                problems.append(f"нет ссылки: [{title}] «{text}» ({cc.coordinate})")
                continue
            posts.append({"title": text, "url": url})
        cats.append({"id": slug, "slug": slug, "emoji": emoji, "title": title,
                     "subtitle": sub, "pain": pain, "posts": posts})
    return {"appId": app_id, "channel": channel, "categories": cats}, problems

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("out")
    ap.add_argument("--app", required=True); ap.add_argument("--channel", required=True)
    a = ap.parse_args()
    data, problems = convert(a.xlsx, a.app, a.channel)
    for p in problems:
        print("ВНИМАНИЕ:", p, file=sys.stderr)
    json.dump(data, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
    n = sum(len(c["posts"]) for c in data["categories"])
    print(f"готово: {len(data['categories'])} категорий, {n} постов -> {a.out}")
    sys.exit(1 if problems else 0)
